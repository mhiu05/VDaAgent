"""Application service for profiling workflows.

The service is the API-facing orchestrator: validate source type, inspect schema,
plan metrics, execute DuckDB profiling, then normalize the result for responses.
"""

from __future__ import annotations

import csv
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

from src.analysis.statistical_tests import StatisticalTestService
from src.models.schemas import (
    DatabaseConnectionConfig,
    DatabaseConnectionStatus,
    DatabasePreviewResult,
    DatabaseQueryPreviewResult,
    DatabaseQueryRequest,
    DatabaseStatisticalTestRequest,
    DatabaseTablesResult,
    DatabaseTableInfo,
    FilePreviewResult,
    ProfileCollectionResult,
    ProfileCollectionSummary,
    ProfileColumnsResult,
    ProfileCorrelationsResult,
    ProfileFindingsResult,
    ProfileMetadata,
    ProfileRelationships,
    ProfileResult,
    ProfileSchemaColumn,
    ProfileSchemaResult,
    ProfileSectionsResult,
    StatisticalTestBatchResult,
    StatisticalTestResult,
    StatisticalTestSpec,
)
from src.profiling.executors.database import DatabaseProfileExecutor
from src.profiling.executors.duckdb import DuckDBProfileExecutor
from src.profiling.planner import ProfilingPlanner
from src.profiling.relationships import infer_relationships
from src.profiling.result_normalizer import ProfileResultNormalizer


class ProfilingService:
    def __init__(self) -> None:
        self.planner = ProfilingPlanner()
        self.normalizer = ProfileResultNormalizer()

    def profile_csv_file(self, file_path: Path, source_name: str) -> ProfileResult:
        executor = DuckDBProfileExecutor()
        schema = executor.get_schema(file_path)
        plan = self.planner.build_plan(schema)
        raw_result = executor.profile_file(file_path, source_name, plan)
        return self.normalizer.normalize(raw_result)

    def profile_csv_files(self, files: list[tuple[Path, str]]) -> ProfileCollectionResult:
        profiles = [self.profile_csv_file(file_path, source_name) for file_path, source_name in files]
        return self._build_collection_result("multi_csv", "uploaded_csv_files", profiles)

    def profile_excel_workbook(self, workbook_path: Path, workbook_name: str, temp_dir: Path) -> ProfileCollectionResult:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel profiling requires `pip install openpyxl`.") from exc

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        csv_files: list[tuple[Path, str]] = []
        for sheet in workbook.worksheets:
            csv_path = temp_dir / f"{_safe_filename(workbook_path.stem)}__{_safe_filename(sheet.title)}.csv"
            self._write_sheet_to_csv(sheet, csv_path)
            csv_files.append((csv_path, f"{workbook_name}:{sheet.title}"))
        workbook.close()

        profiles = [self.profile_csv_file(file_path, source_name) for file_path, source_name in csv_files]
        return self._build_collection_result("excel_workbook", workbook_name, profiles)

    def inspect_csv_schema(self, file_path: Path, source_name: str) -> ProfileSchemaResult:
        executor = DuckDBProfileExecutor()
        inspection = executor.inspect_file(file_path)
        schema = inspection["schema"]
        return ProfileSchemaResult(
            source_name=source_name,
            source_type="file",
            row_count=inspection["row_count"],
            column_count=len(schema),
            columns=[
                ProfileSchemaColumn(name=column["name"], data_type=column["data_type"])
                for column in schema
            ],
        )

    def preview_csv_file(self, file_path: Path, source_name: str, limit: int = 50) -> FilePreviewResult:
        rows = DuckDBProfileExecutor().preview_file(file_path, limit)
        return FilePreviewResult(source_name=source_name, source_type="file", rows=rows)

    def profile_csv_columns(self, file_path: Path, source_name: str) -> ProfileColumnsResult:
        full_result = self.profile_csv_file(file_path, source_name)
        return ProfileColumnsResult(
            source_name=full_result.source.name,
            source_type=full_result.source.type,
            row_count=full_result.dataset_summary.row_count,
            column_count=full_result.dataset_summary.column_count,
            columns=full_result.columns,
        )

    def profile_csv_correlations(self, file_path: Path, source_name: str) -> ProfileCorrelationsResult:
        full_result = self.profile_csv_file(file_path, source_name)
        return ProfileCorrelationsResult(
            source_name=full_result.source.name,
            source_type=full_result.source.type,
            correlations=full_result.relationships.correlations,
        )

    def profile_csv_findings(self, file_path: Path, source_name: str) -> ProfileFindingsResult:
        full_result = self.profile_csv_file(file_path, source_name)
        return ProfileFindingsResult(
            source_name=full_result.source.name,
            source_type=full_result.source.type,
            findings=full_result.findings,
        )

    def profile_csv_sections(
        self,
        file_path: Path,
        source_name: str,
        sections: list[str],
    ) -> ProfileSectionsResult:
        full_result = self.profile_csv_file(file_path, source_name)
        return self._build_sections_result(full_result, sections)

    def run_csv_statistical_test(
        self,
        file_path: Path,
        source_name: str,
        test_type: str,
        alpha: float = 0.05,
        x_column: str | None = None,
        y_column: str | None = None,
        value_column: str | None = None,
        group_column: str | None = None,
    ) -> StatisticalTestResult:
        return StatisticalTestService().run_csv_test(
            file_path=file_path,
            source_name=source_name,
            test_type=test_type,
            alpha=alpha,
            x_column=x_column,
            y_column=y_column,
            value_column=value_column,
            group_column=group_column,
        )

    def run_csv_statistical_tests(
        self,
        file_path: Path,
        source_name: str,
        tests: list[StatisticalTestSpec],
    ) -> StatisticalTestBatchResult:
        results, errors = StatisticalTestService().run_csv_tests(file_path, source_name, tests)
        return StatisticalTestBatchResult(source_name=source_name, results=results, errors=errors)

    def run_database_statistical_test(
        self,
        request: DatabaseStatisticalTestRequest,
    ) -> StatisticalTestResult:
        columns = _required_test_columns(
            request.test_type,
            request.x_column,
            request.y_column,
            request.value_column,
            request.group_column,
        )
        executor = DatabaseProfileExecutor(request.connection)
        table = executor.query_builder.table_reference(request.table_name, request.schema_name)
        selected_columns = ", ".join(executor.query_builder.quote_identifier(column) for column in columns)
        query = f"SELECT {selected_columns} FROM {table}"
        with executor.engine.connect() as connection:
            rows = connection.execute(executor.text(query)).mappings().fetchall()

        source_name = f"{request.schema_name}.{request.table_name}" if request.schema_name else request.table_name
        with TemporaryDirectory() as temp_dir:
            csv_path = Path(temp_dir) / "statistical_test_source.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as file:
                writer = csv.DictWriter(file, fieldnames=columns)
                writer.writeheader()
                for row in rows:
                    writer.writerow({column: row.get(column) for column in columns})
            return self.run_csv_statistical_test(
                file_path=csv_path,
                source_name=source_name,
                test_type=request.test_type,
                alpha=request.alpha,
                x_column=request.x_column,
                y_column=request.y_column,
                value_column=request.value_column,
                group_column=request.group_column,
            )

    def run_csv_pearson_correlation(
        self,
        file_path: Path,
        source_name: str,
        x_column: str,
        y_column: str,
        alpha: float = 0.05,
    ) -> StatisticalTestResult:
        return StatisticalTestService().pearson_correlation(
            file_path, source_name, x_column, y_column, alpha
        )

    def run_csv_spearman_correlation(
        self,
        file_path: Path,
        source_name: str,
        x_column: str,
        y_column: str,
        alpha: float = 0.05,
    ) -> StatisticalTestResult:
        return StatisticalTestService().spearman_correlation(
            file_path, source_name, x_column, y_column, alpha
        )

    def run_csv_independent_t_test(
        self,
        file_path: Path,
        source_name: str,
        value_column: str,
        group_column: str,
        alpha: float = 0.05,
    ) -> StatisticalTestResult:
        return StatisticalTestService().independent_t_test(
            file_path, source_name, value_column, group_column, alpha
        )

    def run_csv_chi_square_independence(
        self,
        file_path: Path,
        source_name: str,
        x_column: str,
        y_column: str,
        alpha: float = 0.05,
    ) -> StatisticalTestResult:
        return StatisticalTestService().chi_square_independence(
            file_path, source_name, x_column, y_column, alpha
        )

    def run_csv_one_way_anova(
        self,
        file_path: Path,
        source_name: str,
        value_column: str,
        group_column: str,
        alpha: float = 0.05,
    ) -> StatisticalTestResult:
        return StatisticalTestService().one_way_anova(
            file_path, source_name, value_column, group_column, alpha
        )

    def test_database_connection(
        self,
        config: DatabaseConnectionConfig,
    ) -> DatabaseConnectionStatus:
        result = DatabaseProfileExecutor(config).test_connection()
        return DatabaseConnectionStatus(**result)

    def list_database_tables(self, config: DatabaseConnectionConfig) -> DatabaseTablesResult:
        executor = DatabaseProfileExecutor(config)
        tables = executor.list_tables()
        return DatabaseTablesResult(
            source_type=config.type,
            database=config.database,
            tables=[DatabaseTableInfo(**table) for table in tables],
        )

    def inspect_database_table(
        self,
        config: DatabaseConnectionConfig,
        table_name: str,
        schema_name: str | None = None,
    ) -> ProfileSchemaResult:
        executor = DatabaseProfileExecutor(config)
        schema = executor.get_schema(table_name, schema_name)
        with executor.engine.connect() as connection:
            row_count = connection.execute(
                executor.text(executor.query_builder.row_count_query(table_name, schema_name))
            ).fetchone()[0]
        source_name = f"{schema_name}.{table_name}" if schema_name else table_name
        return ProfileSchemaResult(
            source_name=source_name,
            source_type=config.type,
            row_count=row_count,
            column_count=len(schema),
            columns=[
                ProfileSchemaColumn(name=column["name"], data_type=column["data_type"])
                for column in schema
            ],
        )

    def preview_database_table(
        self,
        config: DatabaseConnectionConfig,
        table_name: str,
        schema_name: str | None = None,
        limit: int = 20,
    ) -> DatabasePreviewResult:
        rows = DatabaseProfileExecutor(config).preview_table(table_name, schema_name, limit)
        return DatabasePreviewResult(
            source_type=config.type,
            schema_name=schema_name,
            table_name=table_name,
            rows=rows,
        )

    def preview_database_query(self, request: DatabaseQueryRequest) -> DatabaseQueryPreviewResult:
        executor = DatabaseProfileExecutor(request.connection)
        query = _read_only_query(request.query)
        preview_query = _limited_query(query, request.connection.type, request.limit)
        with executor.engine.connect() as connection:
            result = connection.execute(executor.text(preview_query))
            rows = [dict(row) for row in result.mappings().fetchall()]
        columns = _infer_schema_from_rows(rows)
        return DatabaseQueryPreviewResult(
            source_type=request.connection.type,
            query=query,
            row_count=len(rows),
            column_count=len(columns),
            columns=columns,
            rows=rows,
        )

    def profile_database_query(self, request: DatabaseQueryRequest) -> ProfileResult:
        executor = DatabaseProfileExecutor(request.connection)
        query = _read_only_query(request.query)
        with executor.engine.connect() as connection:
            result = connection.execute(executor.text(query))
            rows = [dict(row) for row in result.mappings().fetchall()]
            columns = list(result.keys())

        with TemporaryDirectory() as temp_dir:
            csv_path = Path(temp_dir) / "database_query_result.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as file:
                writer = csv.DictWriter(file, fieldnames=columns)
                writer.writeheader()
                writer.writerows(rows)
            return self.profile_csv_file(csv_path, "database_query")

    def profile_database_table(
        self,
        config: DatabaseConnectionConfig,
        table_name: str,
        schema_name: str | None = None,
    ) -> ProfileResult:
        executor = DatabaseProfileExecutor(config)
        schema = executor.get_schema(table_name, schema_name)
        plan = self.planner.build_plan(schema)
        raw_result = executor.profile_table(table_name, schema_name, plan)
        return self.normalizer.normalize(raw_result)

    def profile_database_sections(
        self,
        config: DatabaseConnectionConfig,
        table_name: str,
        schema_name: str | None,
        sections: list[str],
    ) -> ProfileSectionsResult:
        full_result = self.profile_database_table(config, table_name, schema_name)
        return self._build_sections_result(full_result, sections)

    def _build_sections_result(
        self,
        full_result: ProfileResult,
        sections: list[str],
    ) -> ProfileSectionsResult:
        selected: dict[str, object] = {}
        errors: list[dict[str, str]] = []
        normalized_sections = [section.strip().lower() for section in sections if section.strip()]

        for section in normalized_sections:
            if section == "full":
                selected["full"] = full_result.model_dump()
            elif section == "metadata":
                selected["metadata"] = full_result.profile_metadata.model_dump()
            elif section == "source":
                selected["source"] = full_result.source.model_dump()
            elif section == "summary":
                selected["summary"] = full_result.dataset_summary.model_dump()
            elif section == "schema":
                selected["schema"] = [
                    {"name": column.name, "data_type": column.data_type}
                    for column in full_result.columns
                ]
            elif section == "columns":
                selected["columns"] = [column.model_dump() for column in full_result.columns]
            elif section == "correlations":
                selected["correlations"] = [
                    correlation.model_dump()
                    for correlation in full_result.relationships.correlations
                ]
            elif section == "relationships":
                selected["relationships"] = full_result.relationships.model_dump()
            elif section == "findings":
                selected["findings"] = [finding.model_dump() for finding in full_result.findings]
            elif section == "quality_summary":
                selected["quality_summary"] = full_result.quality_summary.model_dump()
            else:
                errors.append({"section": section, "error": "Unsupported profiling section."})

        return ProfileSectionsResult(
            source_name=full_result.source.name,
            source_type=full_result.source.type,
            requested_sections=normalized_sections,
            sections=selected,
            errors=errors,
        )

    def _build_collection_result(
        self,
        collection_type: str,
        collection_name: str,
        profiles: list[ProfileResult],
    ) -> ProfileCollectionResult:
        return ProfileCollectionResult(
            profile_metadata=ProfileMetadata(
                engine="duckdb",
                profile_mode="collection",
                sampled=False,
                sample_size=None,
                generated_at=datetime.now(timezone(timedelta(hours=7))).isoformat(),
                profiling_version="0.1.0",
                outlier_method="iqr",
                correlation_method="pearson",
                hitl_required=True,
            ),
            collection_type=collection_type,
            collection_name=collection_name,
            collection_summary=ProfileCollectionSummary(
                source_count=len(profiles),
                total_row_count=sum(profile.dataset_summary.row_count for profile in profiles),
                total_column_count=sum(profile.dataset_summary.column_count for profile in profiles),
            ),
            sources=profiles,
            relationships=ProfileRelationships(inferred_relationships=infer_relationships(profiles)),
        )

    def _write_sheet_to_csv(self, sheet, csv_path: Path) -> None:
        with csv_path.open("w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])


def _safe_filename(value: str) -> str:
    safe = "".join(char if char.isalnum() or char in {"-", "_"} else "_" for char in value)
    return safe.strip("_") or "sheet"


def _read_only_query(query: str) -> str:
    cleaned = query.strip().rstrip(";")
    normalized = cleaned.lower()
    if not (normalized.startswith("select ") or normalized.startswith("with ")):
        raise ValueError("Only read-only SELECT queries are supported.")
    blocked_tokens = [" insert ", " update ", " delete ", " drop ", " alter ", " create ", " truncate ", " merge "]
    padded = f" {normalized} "
    if any(token in padded for token in blocked_tokens):
        raise ValueError("Query contains a statement that is not allowed for profiling preview.")
    return cleaned


def _limited_query(query: str, database_type: str, limit: int) -> str:
    if database_type == "sql_server":
        return f"SELECT TOP ({limit}) * FROM ({query}) AS profiling_query_preview"
    return f"SELECT * FROM ({query}) AS profiling_query_preview LIMIT {limit}"


def _infer_schema_from_rows(rows: list[dict[str, object | None]]) -> list[ProfileSchemaColumn]:
    if not rows:
        return []
    columns = []
    for column_name in rows[0].keys():
        value = next((row.get(column_name) for row in rows if row.get(column_name) is not None), None)
        columns.append(ProfileSchemaColumn(name=column_name, data_type=_python_value_type(value)))
    return columns


def _python_value_type(value: object | None) -> str:
    if value is None:
        return "UNKNOWN"
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, int):
        return "BIGINT"
    if isinstance(value, float):
        return "DOUBLE"
    if isinstance(value, datetime):
        return "TIMESTAMP"
    return "VARCHAR"


def _required_test_columns(
    test_type: str,
    x_column: str | None,
    y_column: str | None,
    value_column: str | None,
    group_column: str | None,
) -> list[str]:
    normalized = test_type.lower().replace("-", "_")
    if normalized in {"pearson_correlation", "spearman_correlation", "chi_square_independence"}:
        columns = [x_column, y_column]
    elif normalized in {"independent_t_test", "one_way_anova"}:
        columns = [value_column, group_column]
    else:
        raise ValueError("Unsupported database statistical test type.")
    if any(column is None or not str(column).strip() for column in columns):
        raise ValueError("Selected statistical test requires two column fields.")
    return list(dict.fromkeys(str(column) for column in columns))
